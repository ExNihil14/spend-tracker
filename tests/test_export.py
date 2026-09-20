"""Экспорт транзакций: CSV (utf-8-sig, «;»), XLSX (openpyxl), store-метод без лимита страницы."""
from __future__ import annotations

import csv
import io
from datetime import date

import pytest
from fastapi.testclient import TestClient

from spendtrack import cli
from spendtrack.export import (
    CSV_HEADERS,
    csv_bytes,
    export_filename,
    export_row,
    plain_amount,
    write_xlsx,
)
from spendtrack.store import Store

TX = {
    "date": "2026-09-20",
    "description": 'ПИЦЦА "ДОДО"; МОСКВА',
    "amount_kopecks": -12345,
    "category": "restaurants",
    "category_source": "llm",
    "confidence": 0.98,
    "merchant": "DODO",
    "account_anon": "acc_abcd1234",
    "review_status": "approved",
    "category_llm": "restaurants",
}


def _store_with_txs(tmp_path, name="exp.db") -> Store:
    store = Store(db_path=tmp_path / name)
    for i in range(3):
        store.add_transaction(date=f"2026-09-0{i + 1}", description=f"TX {i}",
                              amount_kopecks=-100 * (i + 1), category="other",
                              category_source="manual")
    return store


def test_plain_amount_ascii_minus_and_kopecks():
    assert plain_amount(-12345) == "-123.45"
    assert plain_amount(0) == "0.00"
    assert plain_amount(5) == "0.05"
    assert plain_amount(250000) == "2500.00"


def test_csv_has_bom_delimiter_and_escaped_fields():
    data = csv_bytes([TX])
    assert data.startswith(b"\xef\xbb\xbf")
    text = data.decode("utf-8-sig")
    rows = list(csv.reader(io.StringIO(text), delimiter=";"))
    assert rows[0] == list(CSV_HEADERS)
    assert rows[1][0] == "2026-09-20"
    assert rows[1][1] == 'ПИЦЦА "ДОДО"; МОСКВА'  # поле с «;» и кавычками экранировано csv-модулем
    assert rows[1][2] == "-123.45"
    assert rows[1][5] == "0.98"
    assert "\r\n" in text  # Excel-совместимые переводы строк


def test_csv_empty_has_header_only():
    text = csv_bytes([]).decode("utf-8-sig")
    assert text.splitlines() == [";".join(CSV_HEADERS)]


def test_export_row_handles_missing_fields():
    row = export_row({"date": "2026-01-01", "description": "X", "amount_kopecks": 100,
                      "category": "other"})
    assert row == ("2026-01-01", "X", "1.00", "other", "", "", "", "", "", "")


def test_formula_injection_is_neutralized():
    """Текстовые поля с ведущими =,+,-,@ получают `'`; сумма остаётся числовой."""
    tx = dict(TX, description="=SUM(A1:A2)", merchant="+79001234567", category_llm="@cmd")
    row = export_row(tx)
    assert row[1] == "'=SUM(A1:A2)"
    assert row[6] == "'+79001234567"
    assert row[9] == "'@cmd"
    assert row[2] == "-123.45"  # суммы не экранируются


def test_export_filename_has_date_and_ext():
    assert export_filename("csv").startswith("spend-export-")
    assert export_filename("csv").endswith(".csv")


def test_write_xlsx_native_types(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "out.xlsx"
    assert write_xlsx([dict(TX, description="=SUM(A1)")], path) == 1
    wb = load_workbook(path)
    ws = wb.active
    assert [c.value for c in ws[1]] == list(CSV_HEADERS)
    assert ws["A2"].value.date() == date(2026, 9, 20)  # Excel отдаёт дату как datetime
    assert ws["C2"].value == -123.45
    assert ws["C2"].number_format == "#,##0.00"
    assert ws["B2"].value == "'=SUM(A1)"  # не формула
    assert ws.auto_filter.ref == "A1:J2"
    assert ws.freeze_panes == "A2"
    wb.close()


def test_write_xlsx_empty(tmp_path):
    """Пустой экспорт — валидный файл с шапкой и freeze, без автофильтра на пустой таблице."""
    from openpyxl import load_workbook

    path = tmp_path / "empty.xlsx"
    assert write_xlsx([], path) == 0
    wb = load_workbook(path)
    ws = wb.active
    assert [c.value for c in ws[1]] == list(CSV_HEADERS)
    assert ws.max_row == 1
    assert ws.auto_filter.ref is None
    assert ws.freeze_panes == "A2"
    wb.close()


def test_write_xlsx_formula_prefixes(tmp_path):
    """XLSX: опасные префиксы в текстовых полях не становятся формулами."""
    from openpyxl import load_workbook

    path = tmp_path / "inj.xlsx"
    write_xlsx([dict(TX, description="+1+1", merchant="-cmd", category_llm="@x")], path)
    wb = load_workbook(path)
    ws = wb.active
    assert ws["B2"].value == "'+1+1"
    assert ws["G2"].value == "'-cmd"
    assert ws["J2"].value == "'@x"
    wb.close()


def test_export_transactions_all_rows_order_and_filters(tmp_path):
    store = _store_with_txs(tmp_path)
    try:
        txs = store.export_transactions()
        assert isinstance(txs, tuple)  # иммутабельная граница
        assert [t["description"] for t in txs] == ["TX 0", "TX 1", "TX 2"]  # хронология
        assert len(store.export_transactions(month="2026-09")) == 3
        assert len(store.export_transactions(month="2026-10")) == 0
        assert len(store.export_transactions(category="other")) == 3
        assert len(store.export_transactions(search="TX 1")) == 1
        assert len(store.export_transactions(date_from="2026-09-02")) == 2
        assert len(store.export_transactions(date_to="2026-09-02")) == 2
        # пересечение фильтров: месяц + диапазон
        assert [t["description"] for t in store.export_transactions(
            month="2026-09", date_from="2026-09-02")] == ["TX 1", "TX 2"]
    finally:
        store.close()


def test_export_transactions_no_page_limit(tmp_path):
    """Экспорт не подчиняется лимиту страницы списка (500) — выгружается всё."""
    store = Store(db_path=tmp_path / "big.db")
    try:
        for i in range(501):
            store.add_transaction(date="2026-09-01", description=f"TX {i:03d}",
                                  amount_kopecks=-100, category="other", category_source="manual")
        assert len(store.export_transactions()) == 501
    finally:
        store.close()


def test_cli_export_csv_and_xlsx(tmp_path, monkeypatch, capsys):
    db = tmp_path / "cli.db"
    monkeypatch.setenv("SPENDTRACK_DB_PATH", str(db))
    store = _store_with_txs(tmp_path, name="cli.db")
    store.close()

    out_csv = tmp_path / "e.csv"
    assert cli.main(["export", "--out", str(out_csv)]) == 0
    assert out_csv.read_bytes().startswith(b"\xef\xbb\xbf")
    assert "OK 3 транзакций" in capsys.readouterr().out

    out_xlsx = tmp_path / "e.xlsx"
    assert cli.main(["export", "--format", "xlsx", "--out", str(out_xlsx)]) == 0
    assert out_xlsx.read_bytes()[:2] == b"PK"


def test_cli_export_rejects_bad_dates(tmp_path, monkeypatch):
    """Неверный формат даты/месяца — явная ошибка argparse, а не тихий пустой файл."""
    monkeypatch.setenv("SPENDTRACK_DB_PATH", str(tmp_path / "bad.db"))
    for argv in (["export", "--from", "2026/09/01"],
                 ["export", "--to", "01-09-2026"],
                 ["export", "--month", "2026-9"]):
        with pytest.raises(SystemExit) as e:
            cli.main(argv)
        assert e.value.code == 2


def test_export_routes_and_link(tmp_path, monkeypatch):
    db = tmp_path / "api.db"
    monkeypatch.setenv("SPENDTRACK_DB_PATH", str(db))
    store = _store_with_txs(tmp_path, name="api.db")
    store.close()

    from spendtrack.main import app

    with TestClient(app) as client:
        page = client.get("/")
        assert page.status_code == 200
        assert "/export.csv" in page.text
        assert 'aria-label="Экспорт' in page.text

        r = client.get("/export.csv", params={"month": "2026-09"})
        assert r.status_code == 200
        assert r.content.startswith(b"\xef\xbb\xbf")
        assert r.headers["content-type"].startswith("text/csv")
        assert "attachment" in r.headers["content-disposition"]
        assert "TX 0" in r.content.decode("utf-8-sig")

        # веб-фильтры category+q работают так же, как у списка
        rc = client.get("/export.csv", params={"category": "other", "q": "TX 1"})
        text = rc.content.decode("utf-8-sig")
        assert "TX 1" in text and "TX 0" not in text

        rx = client.get("/export.xlsx")
        assert rx.status_code == 200
        assert rx.content[:2] == b"PK"
        assert "spreadsheet" in rx.headers["content-type"]
