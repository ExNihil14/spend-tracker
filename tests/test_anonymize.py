"""Anonymizer выписок (P1 #3, K2): PII заменены, формат сохранён, образец — валидная фикстура.

CLI-команда — `spendtrack anonymize` (модуль `spendtrack.anonymize`); даты и суммы не обезличиваются,
по умолчанию остаются первые 5 строк, о чём есть явное предупреждение в stderr.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from synth_bank import gen_sber

from spendtrack.anonymize import DEFAULT_ROWS, anonymize_csv, anonymize_xlsx, main, run_anonymize
from spendtrack.csv_import import import_csv
from spendtrack.store import Store

CSV = (
    "Номер документа;Дата операции;Дата платежа;Номер карты;Статус;Сумма операции;"
    "Валюта операции;Категория;Описание\n"
    "777;01.09.2026 10:00;01.09.2026;4111111111111111;Выполнено;-1234,56;RUB;Супермаркеты;ЛЕНТА ПЯТЁРОЧКА\n"
    "778;02.09.2026 11:00;02.09.2026;4111111111111111;Выполнено;-500,00;RUB;Супермаркеты;ЛЕНТА ПЯТЁРОЧКА\n"
    "779;03.09.2026 12:00;03.09.2026;4111111111111111;Выполнено;-700,00;RUB;Кафе;КАФЕ МОЛОКО\n"
)


def test_pii_removed_and_format_kept():
    out, _report = anonymize_csv(CSV)
    assert "ЛЕНТА" not in out and "КАФЕ" not in out
    assert "4111111111111111" not in out and ";777;" not in out
    # шапка/даты/суммы/статусы/категории — как в оригинале (образец остаётся валидным CSV)
    assert out.splitlines()[0] == CSV.splitlines()[0]
    assert "01.09.2026" in out and "-1234,56" in out and "Супермаркеты" in out and "Выполнено" in out


def test_same_value_same_pseudonym():
    out, _report = anonymize_csv(CSV)
    lines = out.splitlines()[1:]
    assert lines[0].split(";")[-1] == lines[1].split(";")[-1]
    assert lines[0].split(";")[-1] != lines[2].split(";")[-1]


def test_doc_and_account_columns_replaced():
    out, report = anonymize_csv(CSV)
    row = out.splitlines()[1].split(";")
    assert row[0] == "1"                       # номер документа → порядковый
    assert row[3].startswith("КАРТА_")         # номер карты → псевдоним
    assert report["anonymized"]["Описание"] == "desc"
    assert "Номер карты" in report["anonymized"]


def test_extra_column_flag():
    out, _report = anonymize_csv("ФИО;Сумма операции;Описание\nИванов И.И.;-100;ЛЕНТА\n",
                                 {"фио"})
    assert "Иванов" not in out
    assert "ОПЕРАЦИЯ_" in out


def test_cp1251_input_decoded():
    out, _report = anonymize_csv(CSV.encode("cp1251"))
    assert "ЛЕНТА" not in out
    assert "ОПЕРАЦИЯ_0001" in out


def test_untouched_columns_reported():
    _out, report = anonymize_csv(CSV)
    assert "Дата операции" in report["untouched"]
    assert "Статус" in report["untouched"]


def test_empty_input_returns_empty_report():
    """Пустой ввод — пустой результат без исключений (аудит 23.09, P2)."""
    text, report = anonymize_csv(b"")
    assert text == ""
    assert report == {"anonymized": {}, "untouched": [], "empty_headers": [], "unique": {},
                      "rows_total": 0, "rows_written": 0}


def test_rows_limit_keeps_first_rows_and_counts():
    """--rows N: в файле остаются первые N строк данных; даты/суммы не изменяются (K2)."""
    out, report = anonymize_csv(CSV, max_rows=2)
    assert report["rows_total"] == 3 and report["rows_written"] == 2
    lines = out.splitlines()
    assert len(lines) == 3  # шапка + 2 строки
    assert "01.09.2026" in lines[1] and "02.09.2026" in lines[2]
    assert "03.09.2026" not in out


def test_rows_zero_keeps_all():
    out, report = anonymize_csv(CSV, max_rows=0)
    assert report["rows_written"] == 3
    assert "03.09.2026" in out


def test_negative_rows_rejected(tmp_path, capsys):
    """Отрицательный лимит — ошибка, а не тихий «весь файл» (ревью $0, P1)."""
    with pytest.raises(ValueError):
        anonymize_csv(CSV, max_rows=-1)

    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    assert run_anonymize(src, max_rows=-1) == 1
    assert "не может быть отрицательным" in capsys.readouterr().err
    assert not (tmp_path / "s.anon.csv").exists()


def test_cli_output_argument_conflict(tmp_path):
    """OUT и -o одновременно — явная ошибка, а не тихий приоритет одного из них."""
    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    with pytest.raises(SystemExit) as e:
        main([str(src), "a.csv", "-o", "b.csv"])
    assert e.value.code == 2


def test_cli_writes_anon_file_and_reports(tmp_path, capsys):
    """CLI: пишет <имя>.anon.csv, печатает отчёт и предупреждение о датах/суммах (audit 23.09)."""
    src = tmp_path / "statement.csv"
    src.write_text(CSV, encoding="utf-8")

    assert main([str(src)]) == 0
    dst = tmp_path / "statement.anon.csv"
    assert dst.is_file()
    captured = capsys.readouterr()
    assert "Записано:" in captured.out and "Обезличено:" in captured.out
    assert "ВНИМАНИЕ" in captured.err and "даты, суммы" in captured.err

    text = dst.read_text(encoding="utf-8")
    assert "ЛЕНТА ПЯТЁРОЧКА" not in text
    assert "4111111111111111" not in text


def test_cli_custom_out_and_extra_column(tmp_path, capsys):
    """-o и --anon-column: путь настраивается, дополнительная колонка обезличивается."""
    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    custom = tmp_path / "custom.csv"

    assert main([str(src), "-o", str(custom), "--anon-column", "Категория"]) == 0
    text = custom.read_text(encoding="utf-8")
    assert "Супермаркеты" not in text


def test_cli_positional_out_and_rows(tmp_path, capsys):
    """`spendtrack anonymize IN OUT --rows N` — позиционный OUT и обрезка строк."""
    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    custom = tmp_path / "sample.csv"

    assert main([str(src), str(custom), "--rows", "1"]) == 0
    lines = custom.read_text(encoding="utf-8").splitlines()
    assert len(lines) == 2  # шапка + 1 строка
    captured = capsys.readouterr()
    assert "строк 1 из 3" in captured.out
    assert "оставлено: 1" in captured.err


def test_run_anonymize_default_keeps_5_and_warns(tmp_path, capsys):
    """По умолчанию — первые 5 строк: меньше строк в публичном issue (рекомендация запуска)."""
    src = tmp_path / "long.csv"
    rows = "".join(
        f"{i};0{i}.09.2026;0{i}.09.2026;4111111111111111;Выполнено;-{i}00,00;RUB;X;МАГАЗИН {i}\n"
        for i in range(1, 9)
    )
    src.write_text(CSV.splitlines()[0] + "\n" + rows, encoding="utf-8")

    assert run_anonymize(src) == 0
    dst = tmp_path / "long.anon.csv"
    assert len(dst.read_text(encoding="utf-8").splitlines()) == DEFAULT_ROWS + 1
    captured = capsys.readouterr()
    assert f"строк {DEFAULT_ROWS} из 8" in captured.out
    assert f"оставлено: {DEFAULT_ROWS}" in captured.err


def test_run_anonymize_missing_file_reports_error(tmp_path, capsys):
    assert run_anonymize(tmp_path / "nope.csv") == 1
    assert "не удалось прочитать файл" in capsys.readouterr().err


def test_crlf_terminator_not_doubled(tmp_path):
    """CRLF-выписка (cp1251-экспорт банка) сохраняет «\\r\\n», а не превращается в «\\r\\r\\n»."""
    src = tmp_path / "win.csv"
    src.write_bytes(CSV.replace("\n", "\r\n").encode("cp1251"))

    assert run_anonymize(src, max_rows=0) == 0
    raw = (tmp_path / "win.anon.csv").read_bytes()
    assert raw.count(b"\r\n") == 4  # шапка + 3 строки
    body = raw.replace(b"\r\n", b"")
    assert b"\r" not in body and b"\n" not in body
    assert "ЛЕНТА".encode() not in raw


# ---- XLSX (Ф1 BYN: обезличивание образцов банков; адъюдикация 03.10) ----

def _xlsx_bytes(rows: list[list], header: list[str], title: str | None = None) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if title:
        ws.append([title])
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_pii_removed_and_format_kept():
    from openpyxl import load_workbook

    raw = _xlsx_bytes(
        [["777", "01.09.2026", "4111111111111111", "-1234,56", "Супермаркеты", "ЛЕНТА ПЯТЁРОЧКА"]],
        ["Номер документа", "Дата операции", "Номер карты", "Сумма операции", "Категория", "Описание"])
    out, report = anonymize_xlsx(raw, max_rows=0)
    ws = load_workbook(BytesIO(out)).active
    assert ws.cell(1, 1).value == "Номер документа"        # шапка не тронута
    assert ws.cell(2, 1).value == "1"                      # документ → порядковый
    assert str(ws.cell(2, 3).value).startswith("КАРТА_")   # карта → псевдоним
    assert ws.cell(2, 6).value == "ОПЕРАЦИЯ_0001"          # описание → псевдоним
    assert ws.cell(2, 2).value == "01.09.2026"             # дата не тронута
    assert ws.cell(2, 4).value == "-1234,56"               # сумма не тронута
    assert report["rows_total"] == 1 and report["anonymized"]["Описание"] == "desc"


def test_xlsx_title_row_before_header():
    from openpyxl import load_workbook

    raw = _xlsx_bytes([["1", "ЛЕНТА"]], ["Номер документа", "Описание"],
                      title="Выписка по карточке")
    out, _report = anonymize_xlsx(raw, max_rows=0)
    ws = load_workbook(BytesIO(out)).active
    assert ws.cell(1, 1).value == "Выписка по карточке"    # титульная строка сохранена
    assert ws.cell(2, 2).value == "Описание"               # шапка найдена ниже титула
    assert ws.cell(3, 2).value == "ОПЕРАЦИЯ_0001"


def test_xlsx_rows_limit_truncates():
    from openpyxl import load_workbook

    raw = _xlsx_bytes([[str(i), f"ЛЕНТА {i}"] for i in range(1, 8)],
                      ["Номер документа", "Описание"])
    out, report = anonymize_xlsx(raw, max_rows=2)
    ws = load_workbook(BytesIO(out)).active
    assert report["rows_total"] == 7 and report["rows_written"] == 2
    assert ws.max_row == 3  # шапка + 2 строки (PII за лимитом удалена, а не «оставлена как есть»)


def test_run_anonymize_xlsx_roundtrip(tmp_path):
    from openpyxl import load_workbook

    src = tmp_path / "v.xlsx"
    src.write_bytes(_xlsx_bytes([["1", "ЛЕНТА"]], ["Номер документа", "Описание"]))
    assert run_anonymize(src, max_rows=0) == 0
    dst = tmp_path / "v.anon.xlsx"
    assert dst.is_file()
    ws = load_workbook(dst).active
    assert ws.cell(2, 2).value == "ОПЕРАЦИЯ_0001"


def test_run_anonymize_legacy_xls_asks_conversion(tmp_path, capsys):
    """Legacy .xls (BIFF) не парсим (NO-GO: без xlrd в core) — просим конвертацию."""
    src = tmp_path / "v.xls"
    src.write_bytes(b"\xd0\xcf\x11\xe0garbage")
    assert run_anonymize(src) == 1
    assert "xlsx" in capsys.readouterr().err.lower()


def test_xlsx_nothing_anonymized_warns(tmp_path, capsys):
    """Ревью wave5 C1: нулевое обезличивание — громкое предупреждение, а не тихий «успех»."""
    src = tmp_path / "v.xlsx"
    src.write_bytes(_xlsx_bytes([["1", "x"]], ["A", "B"]))
    assert run_anonymize(src, max_rows=0) == 0
    assert "НИЧЕГО не обезличено" in capsys.readouterr().err


def test_xlsx_nonstring_cells_ok():
    """Ревью wave5 C2: числа/None в PII-колонках не роняют обезличивание."""
    from openpyxl import load_workbook

    raw = _xlsx_bytes([[4111111111111111, None, 123.45]],
                      ["Номер карты", "Описание", "Сумма операции"])
    out, report = anonymize_xlsx(raw, max_rows=0)
    ws = load_workbook(BytesIO(out)).active
    assert str(ws.cell(2, 1).value).startswith("КАРТА_")
    assert ws.cell(2, 2).value is None  # пустое не трогаем
    assert ws.cell(2, 3).value == 123.45  # сумма не тронута
    assert report["anonymized"]["Номер карты"] == "account"


def test_xlsx_properties_cleared():
    """Ревью wave5 C3/S2: метаданные книги (автор/заголовок) — PII-канал, чистим."""
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    wb.properties.creator = "Иван Иванов"
    wb.properties.title = "Выписка"
    ws = wb.active
    ws.append(["Описание"])
    ws.append(["ЛЕНТА"])
    buf = BytesIO()
    wb.save(buf)
    out, _report = anonymize_xlsx(buf.getvalue(), max_rows=0)
    props = load_workbook(BytesIO(out)).properties
    assert props.creator in ("", None) and props.title in ("", None)


def test_xlsx_above_header_warned(tmp_path, capsys):
    """Ревью wave5 C3: строки выше шапки (титул с ФИО) — предупреждение."""
    src = tmp_path / "t.xlsx"
    src.write_bytes(_xlsx_bytes([["1", "ЛЕНТА"]], ["Номер документа", "Описание"],
                                title="Выписка по карточке Иванова И.И."))
    assert run_anonymize(src, max_rows=0) == 0
    assert "ВНЕ таблицы" in capsys.readouterr().err


def test_run_anonymize_xlsm_rejected(tmp_path, capsys):
    """Ревью wave5 S6: макрос-книга не может быть «обезличена» (VBA теряется) — отказ."""
    src = tmp_path / "v.xlsm"
    src.write_bytes(b"PK\x03\x04garbage")
    assert run_anonymize(src) == 1
    assert "xlsx" in capsys.readouterr().err.lower()


def test_xlsx_phantom_rows_ignored():
    """Wave6 S3: «раздутый» max_row (стили/остатки фильтра) не раздувает отчёт и не вешает delete_rows."""
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Описание", "Сумма"])
    ws.append(["ЛЕНТА", -100])
    ws.cell(50_000, 1)  # обращение к далёкой ячейке раздувает max_row (значения нет)
    buf = BytesIO()
    wb.save(buf)
    out, report = anonymize_xlsx(buf.getvalue(), max_rows=0)
    assert report["rows_total"] == 1  # фантомные строки не считаются
    assert load_workbook(BytesIO(out)).active.max_row >= 2  # файл открывается


def test_xlsx_merge_and_filter_below_cut_dropped():
    """Wave6 S4: merge/автофильтр ниже усечённой части не «висят» за пределами листа."""
    from openpyxl import Workbook, load_workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Описание", "Сумма"])
    for i in range(3):
        ws.append([f"ОПЕРАЦИЯ {i}", -100 - i])
    ws.merge_cells("A6:B6")
    ws.cell(6, 1).value = "Итого"
    ws.auto_filter.ref = "A1:B6"
    buf = BytesIO()
    wb.save(buf)
    out, report = anonymize_xlsx(buf.getvalue(), max_rows=1)
    ws2 = load_workbook(BytesIO(out)).active
    assert report["merged_dropped"] == 1
    assert not ws2.merged_cells.ranges
    assert ws2.auto_filter.ref is None


def test_xlsx_named_csv_routed_by_magic(tmp_path):
    """Wave6 S5: XLSX под именем .csv — роутинг по магии; выход с расширением .xlsx."""
    src = tmp_path / "v.csv"
    src.write_bytes(_xlsx_bytes([["ЛЕНТА", "-100"]], ["Описание", "Сумма"]))
    assert run_anonymize(src, max_rows=0) == 0
    out = src.with_name("v.anon.xlsx")
    assert out.exists() and out.read_bytes()[:4] == b"PK\x03\x04"


def test_html_named_xlsx_asks_conversion(tmp_path, capsys):
    """Wave6 S5: HTML-таблица под именем .xlsx — понятная просьба о конвертации, не «File is not a zip»."""
    src = tmp_path / "v.xlsx"
    src.write_bytes("<html><table><tr><td>ЛЕНТА</td></tr></table>".encode())
    assert run_anonymize(src) == 1
    assert "HTML" in capsys.readouterr().err


def test_anonymized_sample_imports_same_shape(tmp_path):
    """Образец — валидная фикстура: импортируется (sber) с теми же датами/суммами/статусами."""
    raw = gen_sber(n=8, seed=11)
    anon, _report = anonymize_csv(raw)

    def stub(tx, store, taxonomy):
        return {"category": "other", "confidence": 0.5, "merchant": None,
                "source": "llm_pending_review"}

    original = Store(db_path=tmp_path / "a.db")
    sample = Store(db_path=tmp_path / "b.db")
    res_orig = import_csv(raw, original, classify=stub)
    res_anon = import_csv(anon, sample, classify=stub)
    assert res_orig["status"] == res_anon["status"] == "ok"
    assert res_orig["added"] == res_anon["added"]
    assert res_orig["skipped"] == res_anon["skipped"]

    def key(t):
        return (t["date"], t["amount_kopecks"])
    assert (sorted(map(key, original.list_transactions(limit=1000)))
            == sorted(map(key, sample.list_transactions(limit=1000))))
    original.close()
    sample.close()


# ── §J-2: позиционное сохранение формата (csv.reader/writer, пустые/дублирующиеся шапки) ──

def test_trailing_delimiter_and_empty_header_preserved():
    """Хвостовой `;` и колонка без заголовка не теряются (раньше DictWriter выбрасывал пустую шапку)."""
    out, report = anonymize_csv("Дата;Описание;Сумма;\n01.09.2026;ЛЕНТА;-100;\n")
    assert out.splitlines()[0] == "Дата;Описание;Сумма;"
    assert "ЛЕНТА" not in out
    assert report["empty_headers"] == [3]


def test_duplicate_headers_both_anonymized():
    """Дублирующиеся заголовки: DictReader схлопывал колонки — теперь каждая обезличена позиционно."""
    out, _report = anonymize_csv("Описание;Сумма;Описание\nЛЕНТА;-100;МАГНИТ\n")
    cells = out.splitlines()[1].split(";")
    assert cells[0].startswith("ОПЕРАЦИЯ_") and cells[2].startswith("ОПЕРАЦИЯ_")
    assert cells[0] != cells[2]  # разные значения → разные псевдонимы
    assert cells[1] == "-100"
    assert "ЛЕНТА" not in out and "МАГНИТ" not in out


def test_extra_fields_beyond_header_preserved():
    """Поля сверх шапки не отбрасываются (DictReader клал их в restkey и терял)."""
    out, _report = anonymize_csv("Дата;Описание\n01.09.2026;ЛЕНТА;ХВОСТ\n")
    assert out.splitlines()[1].endswith(";ХВОСТ")


def test_missing_anon_column_is_error(tmp_path, capsys):
    """Опечатка в --anon-column не публикует PII молча: ValueError + rc 1, файл не пишется."""
    with pytest.raises(ValueError, match="колонка не найдена"):
        anonymize_csv(CSV, {"неттакой"})

    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    assert run_anonymize(src, extra_columns=("Ф.И.О.",)) == 1
    assert "колонка не найдена" in capsys.readouterr().err
    assert not (tmp_path / "s.anon.csv").exists()


def test_dst_equal_src_rejected(tmp_path, capsys):
    """Перезапись исходной выписки запрещена (безвозвратная потеря данных пользователя)."""
    src = tmp_path / "s.csv"
    src.write_text(CSV, encoding="utf-8")
    assert run_anonymize(src, src) == 1
    assert "совпадают" in capsys.readouterr().err
    assert src.read_text(encoding="utf-8") == CSV
