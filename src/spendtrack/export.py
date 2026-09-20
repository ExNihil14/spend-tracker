"""Экспорт транзакций: CSV (Excel-совместимый, utf-8-sig, «;») и XLSX (openpyxl).

Read-only: функции принимают снимок строк (tuple[dict]) и не трогают Store.
Суммы в CSV — обычное число «-123.45» (ASCII-минус, точка), в XLSX — число с форматом.
"""
from __future__ import annotations

import csv
import io
from collections.abc import Iterable
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

CSV_HEADERS = ("Дата", "Описание", "Сумма", "Категория", "Источник", "Уверенность",
               "Мерчант", "Счёт", "Статус", "Предложение LLM")


def plain_amount(kopecks: int) -> str:
    """Копейки → «-123.45» (ASCII-минус/точка): Excel, Sheets и машинный разбор без магии."""
    sign = "-" if kopecks < 0 else ""
    whole, rest = divmod(abs(kopecks), 100)
    return f"{sign}{whole}.{rest:02d}"


def _confidence(value: Any) -> str:
    try:
        return f"{float(value):.2f}"
    except (TypeError, ValueError):
        return ""


_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value: str) -> str:
    """CSV/formula-инъекция: Excel/Sheets не должны исполнять текст (описание, мерчант, LLM-предложение).

    Опасные ведущие символы получают текстовый префикс `'` (Excel скрывает его при вводе, в CSV — виден).
    Ведущие пробелы не экранируем: Excel их перед `=` не исполняет (проверено), а `strip()` исказил бы данные.
    Суммы не трогаем: они формируются как строгие числа (`plain_amount`) и остаются числовыми.
    """
    return "'" + value if value[:1] in _FORMULA_PREFIXES else value


def export_values(tx: dict) -> tuple:
    """Строка с нативными типами (для XLSX): date, float, защищённые строки."""
    return (
        date.fromisoformat(tx["date"]),
        _safe_text(tx["description"]),
        tx["amount_kopecks"] / 100,
        _safe_text(tx["category"]),
        _safe_text(tx.get("category_source") or ""),
        float(tx.get("confidence") or 0.0),
        _safe_text(tx.get("merchant") or ""),
        _safe_text(tx.get("account_anon") or ""),
        _safe_text(tx.get("review_status") or ""),
        _safe_text(tx.get("category_llm") or ""),
    )


def export_row(tx: dict) -> tuple[str, ...]:
    """Одна строка CSV (порядок = CSV_HEADERS); внутренние поля (fingerprint и пр.) не выгружаем."""
    v = export_values(tx)
    return (
        v[0].isoformat(),
        v[1],
        plain_amount(tx["amount_kopecks"]),
        v[3],
        v[4],
        _confidence(tx.get("confidence")),
        v[6],
        v[7],
        v[8],
        v[9],
    )


def csv_bytes(txs: Iterable[dict]) -> bytes:
    """CSV с BOM (Excel видит UTF-8 кириллицу) и «;» (разделитель RU-локали Excel)."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, delimiter=";", lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(CSV_HEADERS)
    for tx in txs:
        writer.writerow(export_row(tx))
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def write_xlsx(txs: Iterable[dict], path_or_stream: str | Path | io.BytesIO) -> int:
    """XLSX с нативными типами (дата/число), жирной шапкой и автофильтром. Возвращает число строк.

    `txs` итерируется один раз (истощается): для повторных проходов передавайте tuple/list.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Транзакции"
    ws.append(list(CSV_HEADERS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for tx in txs:
        ws.append(list(export_values(tx)))
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0.00"
    widths = (12, 42, 12, 16, 20, 12, 24, 12, 14, 20)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(CSV_HEADERS)).column_letter}{ws.max_row}"
    ws.freeze_panes = "A2"
    wb.save(path_or_stream)
    return ws.max_row - 1


def export_filename(ext: str) -> str:
    """spend-export-YYYYMMDD.<ext> — единое имя для CLI и веб-скачивания."""
    return f"spend-export-{datetime.now(tz=UTC).date().strftime('%Y%m%d')}.{ext}"
