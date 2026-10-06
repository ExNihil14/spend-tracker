"""Обезличить банковскую выписку (CSV/XLSX) для безопасной отправки образца в issue.

Формат сохраняется (шапка/колонки/разделитель/даты/суммы/статусы) — образец остаётся валидной
фикстурой для адаптеров импорта. Заменяются потенциальные PII:
  * описания/мерчанты/назначения → ОПЕРАЦИЯ_0001, … (одинаковые значения → одинаковый псевдоним);
  * номера карт/счетов → КАРТА_0001, …;
  * номера документов/операций → 1, 2, 3, … (в порядке появления).

ВАЖНО: даты, суммы, категории и статусы НЕ обезличиваются. Issue на GitHub публичный — просмотрите
файл перед отправкой. По умолчанию оставляются первые 5 строк данных (`--rows 0` — весь файл).

CLI:  spendtrack anonymize выписка.csv|xlsx [образец] [--rows N] [--anon-column "ФИО"]
Dev:  python -m spendtrack.anonymize … / scripts/anonymize.py (тонкая обёртка; для фикстур — --rows 0)

Legacy `.xls` (BIFF) не поддерживается (NO-GO: без xlrd в core) — сконвертируйте в .xlsx/CSV.
Колонки, не распознанные как PII, скрипт не трогает и перечисляет в отчёте — если в них есть
личные данные (ФИО, адрес), обезличьте их флагом `--anon-column` или удалите колонку.
"""
from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from io import BytesIO, StringIO
from pathlib import Path

from spendtrack.console import utf8_stdout

DESC_ALIASES = {"описание", "назначение", "комментарий", "description", "title", "merchant"}
ACCOUNT_ALIASES = {"номер карты", "номер счёта", "номер счета", "карта", "счёт", "счет", "account"}
DOC_ALIASES = {"номер документа", "номер операции", "номер", "document", "doc", "id"}

DESC_PREFIX = "ОПЕРАЦИЯ"
ACCOUNT_PREFIX = "КАРТА"
KIND_DESC, KIND_ACCOUNT, KIND_DOC = "desc", "account", "doc"
DEFAULT_ROWS = 5


def _strip(value: object) -> str:
    return "" if value is None else str(value).strip()


class _Pseudonyms:
    """Стабильные псевдонимы: одинаковое значение → одинаковый псевдоним (дедуп-фикстуры)."""

    def __init__(self) -> None:
        self._map: dict[tuple[str, str], str] = {}
        self._counters: dict[str, int] = {}

    def get(self, kind: str, value: str) -> str:
        key = (kind, value)
        if key not in self._map:
            self._counters[kind] = self._counters.get(kind, 0) + 1
            n = self._counters[kind]
            if kind == KIND_DOC:
                self._map[key] = str(n)
            else:
                prefix = DESC_PREFIX if kind == KIND_DESC else ACCOUNT_PREFIX
                self._map[key] = f"{prefix}_{n:04d}"
        return self._map[key]

    @property
    def unique(self) -> dict[str, int]:
        """Сколько уникальных значений заменено по каждому типу."""
        return {kind: sum(1 for k in self._map if k[0] == kind)
                for kind in sorted({k[0] for k in self._map})}


def _classify(header: str, extra: set[str]) -> str | None:
    name = header.strip().lower()
    if name in extra:
        return KIND_DESC
    if name in DESC_ALIASES:
        return KIND_DESC
    if name in ACCOUNT_ALIASES:
        return KIND_ACCOUNT
    if name in DOC_ALIASES:
        return KIND_DOC
    return None


def anonymize_csv(  # noqa: C901 — см. cc_ratchet.py
    raw: str | bytes,
    extra_columns: set[str] | None = None,
    max_rows: int | None = None,
) -> tuple[str, dict]:
    """→ (обезличенный CSV, отчёт). `extra_columns` — имена колонок, которые тоже обезличить.

    Формат сохраняется ПОЗИЦИОННО (csv.reader/writer по индексам, §J-2): пустые и дублирующиеся
    заголовки, хвостовой разделитель и лишние поля строк не теряются — образец остаётся валидной
    фикстурой для адаптеров. Неизвестная колонка из `extra_columns` → ValueError (нельзя молча
    опубликовать PII). `max_rows` — оставить первые N строк данных (0/None — все).
    """
    if isinstance(raw, bytes):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1251", errors="replace")
    else:
        text = raw.lstrip("\ufeff")
    if max_rows is not None and max_rows < 0:
        raise ValueError("max_rows не может быть отрицательным (0 — все строки)")
    # csv-парсер падает на одиночном `\r` в незакавыченном поле — нормализуем перевод строки
    # (см. csv_import.import_csv; property-тест §G ловил падение на CR-выгрузках).
    # Терминатор вывода фиксируем ДО нормализации, чтобы CRLF-выписки остались CRLF.
    crlf = "\r\n" in text
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.splitlines()
    if not lines:
        return "", {"anonymized": {}, "untouched": [], "empty_headers": [], "unique": {},
                    "rows_total": 0, "rows_written": 0}

    delimiter = ";" if lines[0].count(";") >= lines[0].count(",") else ","
    reader = csv.reader(StringIO(text), delimiter=delimiter)
    header = next(reader, None)
    if header is None:
        return "", {"anonymized": {}, "untouched": [], "empty_headers": [], "unique": {},
                    "rows_total": 0, "rows_written": 0}
    extra = {c.strip().lower() for c in (extra_columns or set())}
    known = {h.strip().lower() for h in header}
    missing = sorted(extra - known)
    if missing:
        raise ValueError("колонка не найдена: " + ", ".join(missing))
    kinds = [_classify(name, extra) for name in header]
    pseudonyms = _Pseudonyms()
    limit = max_rows if max_rows and max_rows > 0 else None

    out_rows: list[list[str]] = []
    total = 0
    for row in reader:
        total += 1
        if limit is not None and total > limit:
            continue
        new_row = list(row)
        for i, kind in enumerate(kinds):
            if kind is None or i >= len(new_row):
                continue  # короткая строка/лишние поля — сохраняем как есть, не выдумываем колонки
            value = _strip(new_row[i])
            if value:
                new_row[i] = pseudonyms.get(kind, value)
        out_rows.append(new_row)

    terminator = "\r\n" if crlf else "\n"
    buf = StringIO()
    writer = csv.writer(buf, delimiter=delimiter, lineterminator=terminator)
    writer.writerow(header)
    writer.writerows(out_rows)

    report = {
        "anonymized": {name: kind for name, kind in zip(header, kinds) if kind},
        "untouched": [name for name, kind in zip(header, kinds) if not kind and name.strip()],
        "empty_headers": [i for i, name in enumerate(header) if not name.strip()],
        "unique": pseudonyms.unique,
        "rows_total": total,
        "rows_written": len(out_rows),
    }
    return buf.getvalue(), report


def _looks_like_data(value: object) -> bool:
    """Ячейка похожа на ЗНАЧЕНИЕ (а не заголовок): длинные числа, даты, суммы.

    Нужна, чтобы титульная пара «Номер карты | 4111…» не выигрывала у настоящей шапки ниже (C1).
    """
    s = _strip(value)
    if not s:
        return False
    if re.fullmatch(r"[+-]?\d{6,}", s):
        return True
    if re.fullmatch(r"[+-]?\d+[.,]\d{2}", s):
        return True
    return bool(re.fullmatch(r"\d{2}[./-]\d{2}[./-]\d{4}([ T].*)?", s)
                or re.fullmatch(r"\d{4}-\d{2}-\d{2}([ T].*)?", s))


def _find_header(ws, extra: set[str], scan: int = 10) -> tuple[int | None, list | None]:
    """Выбрать строку шапки: оценка всех кандидатов в первых `scan` строках + отказ при неоднозначности.

    Кандидат — строка с хотя бы одной известной колонкой (включая одноколоночные таблицы); больше
    распознанных колонок — сильнее, значения-«данные» среди прочих ячеек — слабее (титул банка).
    Известных колонок нет вовсе — ОТКАЗ (волна 5, окно 06.10): прежний fallback «первая строка с
    ≥2 непустыми ячейками» мог объявить шапкой строку ДАННЫХ и молча записать файл без единого
    обезличенного значения (C2-дыра). Отказ — наверх, вызывающий сам решает (сообщение с --anon-column).
    Ничья между кандидатами → ValueError: тихая догадка могла бы обезличить не ту таблицу (C1).
    """
    ranked: list[tuple[int, int, int, list]] = []
    for r_idx in range(1, min(ws.max_row, scan) + 1):
        values = [ws.cell(r_idx, c).value for c in range(1, ws.max_column + 1)]
        nonempty = [v for v in values if _strip(v)]
        if not nonempty:
            continue
        score = sum(1 for v in values if _strip(v) and _classify(_strip(v), extra))
        suspicious = sum(1 for v in values
                         if _strip(v) and not _classify(_strip(v), extra) and _looks_like_data(v))
        if score:
            ranked.append((score, -suspicious, r_idx, values))
    if not ranked:
        return None, None
    ranked.sort(key=lambda c: (c[0], c[1]), reverse=True)
    best = ranked[0]
    tied = [c[2] for c in ranked if (c[0], c[1]) == (best[0], best[1])]
    if len(tied) > 1:
        raise ValueError(
            "не удалось однозначно выбрать строку-шапку (кандидаты: строки "
            + ", ".join(str(r) for r in tied)
            + ") — приведите лист к одной таблице или уберите лишние заголовочные строки")
    return best[2], best[3]


def _label_columns(header: list, kinds: list, anonymized: dict, untouched: list) -> None:
    for name, kind in zip(header, kinds):
        label = _strip(name)
        if not label:
            continue
        if kind:
            anonymized.setdefault(label, kind)
        else:
            untouched.append(label)


def _last_data_row(ws, header_row: int) -> int:
    """Реальная последняя строка с данными: `max_row` бывает раздут стилями/остатками автофильтра.

    wave6 S3: максимум по УЖЕ существующим ячейкам (`ws._cells`) — обращение через `ws.cell()`
    материализовало бы пустые строки, и `delete_rows` мог получить диапазон в сотни тысяч строк.
    """
    last = header_row
    for (r, _c), cell in ws._cells.items():
        if r > last and _strip(cell.value):
            last = r
    return last


def _truncate_sheet(ws, drop_after: int) -> int:
    """Удалить ВСЁ ниже `drop_after` разреженно; → сколько merged-диапазонов пришлось снять.

    wave6 S4: openpyxl не сдвигает merged-диапазоны/автофильтр при удалении — иначе merge
    «висит» за пределами листа и Excel предлагает восстановление файла.
    anon-7 (wave5): `delete_rows` в openpyxl 3.1.5 материализует весь хвост (`_move_cells` →
    `iter_rows(min_row)` до `max_row`: styled-фантом на строке 1M дал замер ~120 с на 10 строках).
    Удаляем существующие ячейки и размеры строк ниже границы точечно — без движения фантомов.
    """
    dropped = 0
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row > drop_after:
            ws.unmerge_cells(str(mr))
            dropped += 1
    if ws.auto_filter.ref:
        ws.auto_filter.ref = None
    for key in [k for k in list(ws._cells) if k[0] > drop_after]:
        del ws._cells[key]
    for r in [r for r in list(ws.row_dimensions) if r > drop_after]:
        del ws.row_dimensions[r]
    ws._current_row = max((r for r, _ in ws._cells), default=0)
    return dropped


def _count_outside(ws, header_row: int) -> int:
    """Строки ВЫШЕ шапки (титул банка — часто с ФИО) не обезличиваются: считаем и предупреждаем (C3)."""
    return sum(1 for r in range(1, header_row)
               if any(_strip(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))


def _strip_cell_metadata(ws) -> None:
    """should-4 (wave5) + волна 5 (окно 06.10): комментарии/гиперссылки — PII-канал.

    Снимаем со ВСЕХ существующих ячеек; если ЗНАЧЕНИЕ ячейки буквально повторяет цель ссылки
    (Excel хранит URL как текст пустой ячейки), значение тоже очищается — иначе URL (часто с
    токеном) оставался бы в файле после снятия самой гиперссылки.
    """
    for cell in ws._cells.values():
        if cell.comment is not None:
            cell.comment = None
        if cell.hyperlink is not None:
            link = cell.hyperlink
            targets = {t for t in (_strip(getattr(link, "target", None)),
                                   _strip(getattr(link, "location", None))) if t}
            cell.hyperlink = None
            if isinstance(cell.value, str) and _strip(cell.value) in targets:
                cell.value = None


def _check_formulas(ws) -> None:
    """6-lite (wave5): формулы ломают гарантии образца (PII внутри формулы, cached values) — отказ."""
    for (r, c) in sorted(ws._cells):
        cell = ws._cells[(r, c)]
        if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
            raise ValueError(
                f"лист «{ws.title}»: ячейка {cell.coordinate} содержит формулу — обезличивание образца "
                "с формулами не поддерживается; экспортируйте значения (CSV/«только значения») и повторите")


def _anonymize_sheet_rows(ws, header_row: int, kinds: list, keep: list, pseudonyms: _Pseudonyms) -> int:
    """Заменить PII в перечисленных строках листа; → сколько строк записано."""
    written = 0
    for r_idx in keep:
        for c_idx, kind in enumerate(kinds, start=1):
            if kind is None:
                continue
            cell = ws.cell(r_idx, c_idx)
            value = _strip(cell.value)
            if value:
                cell.value = pseudonyms.get(kind, value)
        written += 1
    return written


def anonymize_xlsx(
    raw: bytes,
    extra_columns: set[str] | None = None,
    max_rows: int | None = None,
) -> tuple[bytes, dict]:
    """→ (обезличенный XLSX, отчёт). Контракт — как у `anonymize_csv`.

    Книга правится in-place (openpyxl): листы/форматирование сохраняются, заменяются только значения
    PII-колонок. Шапка ищется в первых 10 строках (у банков бывает титульная строка). Строки данных
    сверх `max_rows` УДАЛЯЮТСЯ — PII не остаётся за лимитом (в отличие от «оставить как есть»).
    """
    from openpyxl import load_workbook

    if max_rows is not None and max_rows < 0:
        raise ValueError("max_rows не может быть отрицательным (0 — все строки)")
    wb = load_workbook(BytesIO(raw))
    extra = {c.strip().lower() for c in (extra_columns or set())}
    pseudonyms = _Pseudonyms()
    anonymized: dict[str, str] = {}
    untouched: list[str] = []
    empty_headers: list[int] = []
    total = written = outside = merged_dropped = 0
    remaining = max_rows if max_rows and max_rows > 0 else None

    for ws in wb.worksheets:
        _strip_cell_metadata(ws)  # should-4: до усечения — в т.ч. пустые ячейки за лимитом
        _check_formulas(ws)
        if not any(_strip(cell.value) for cell in ws._cells.values()):
            continue  # действительно пустой лист — пропускаем (C2)
        header_row, header = _find_header(ws, extra)
        if header_row is None:
            # C2 (wave5) + волна 5: раньше непустой лист без шапки молча сохранялся целиком,
            # а fallback «шапка из строки данных» давал файл без единого обезличенного значения
            raise ValueError(
                f"лист «{ws.title}»: не найдена строка-шапка (нет ни одной известной колонки) — "
                "обезличивание отменено, файл не записан; приведите лист к одной таблице или "
                "укажите --anon-column")
        outside += _count_outside(ws, header_row)
        kinds = [_classify(_strip(name), extra) for name in header]
        empty_headers.extend(i for i, name in enumerate(header) if not _strip(name))
        _label_columns(header, kinds, anonymized, untouched)
        last_data = _last_data_row(ws, header_row)
        rows_all = list(range(header_row + 1, last_data + 1))
        keep = rows_all if remaining is None else rows_all[:remaining]
        total += len(rows_all)
        written += _anonymize_sheet_rows(ws, header_row, kinds, keep, pseudonyms)
        if remaining is not None:
            remaining -= len(keep)
        cut = len(rows_all) - len(keep)
        if cut > 0:
            merged_dropped += _truncate_sheet(ws, header_row + len(keep))

    # метаданные книги — PII-канал: заменяем ВЕСЬ набор на безопасный (allowlist, wave5 should-5;
    # прежний список полей не закрывал subject/identifier и т.п.)
    from openpyxl.packaging.core import DocumentProperties

    props = DocumentProperties()
    props.creator = ""
    wb.properties = props
    buf = BytesIO()
    wb.save(buf)
    report = {
        "anonymized": anonymized,
        "untouched": untouched,
        "empty_headers": empty_headers,
        "unique": pseudonyms.unique,
        "rows_total": total,
        "rows_written": written,
        "outside_rows": outside,
        "merged_dropped": merged_dropped,
        "xlsx": True,
    }
    return buf.getvalue(), report


def _anonymize_payload(raw: bytes, suffix: str, extra: set[str], max_rows: int) -> tuple[bytes, dict]:
    """Роутинг по МАГИИ (расширение — подсказка): ZIP → XLSX; OLE/HTML → понятная просьба; иначе CSV."""
    if suffix == ".xlsm":
        # NO-GO адъюдикации 03.10 + ревью wave5 S6: макрос-книги не тянем (VBA теряется молча)
        raise ValueError(
            "macro .xlsm не поддерживается — сконвертируйте выписку в .xlsx (Excel/LibreOffice) или CSV")
    if raw[:4] == b"PK\x03\x04":
        # wave6 S5: XLSX под именем .csv (частый экспорт) — тоже XLSX, а не «обезличенный» мусор
        return anonymize_xlsx(raw, extra, max_rows=max_rows)
    if raw[:4] == b"\xd0\xcf\x11\xe0":
        raise ValueError(
            "legacy .xls (BIFF/OLE) не поддерживается — сконвертируйте выписку в .xlsx или CSV")
    head = raw[:512].lstrip(b"\xef\xbb\xbf \t\r\n").lower()
    if head.startswith((b"<html", b"<tabl", b"<?xml")):
        raise ValueError("файл похож на HTML/XML-выгрузку — сконвертируйте её в .xlsx или CSV")
    if suffix == ".xlsx":
        raise ValueError(
            "файл с расширением .xlsx не является xlsx-книгой (не zip-контейнер) — сконвертируйте его")
    text, report = anonymize_csv(raw, extra, max_rows=max_rows)
    # bytes, а не write_text: терминатор строк выписки (CRLF в cp1251-файлах) сохраняется
    # как есть — text-mode на Windows превратил бы «\r\n» в «\r\r\n».
    return text.encode("utf-8"), report


def _read_source(src: Path) -> bytes | None:
    try:
        return src.read_bytes()
    except OSError as e:
        print(f"не удалось прочитать файл: {e}", file=sys.stderr, flush=True)
        return None


def _write_anonymized(data: bytes, dst: Path) -> bool:
    try:
        dst.write_bytes(data)
    except OSError as e:
        print(f"не удалось записать файл: {e}", file=sys.stderr, flush=True)
        return False
    return True


def _print_report(report: dict, dst: Path) -> None:
    """Отчёт и предупреждения (stdout/stderr) — вынесено ради cc-лимита run_anonymize."""
    total, written = report["rows_total"], report["rows_written"]
    print(f"Записано: {dst} (строк {written} из {total})")
    if report["anonymized"]:
        print("Обезличено: " + ", ".join(report["anonymized"]))
    if report["untouched"]:
        print("Без изменений (проверьте, нет ли личных данных): " + ", ".join(report["untouched"]))
    warn = [
        ("ВНИМАНИЕ: даты, суммы, категории и статусы не обезличиваются — просмотрите файл "
         "перед отправкой в публичный issue."),
    ]
    if written < total:
        warn.append(f"Строк в файле: {total}, оставлено: {written} (--rows 0 — все).")
    if not report["anonymized"]:
        warn.append("ВНИМАНИЕ: НИЧЕГО не обезличено — колонки не распознаны. Проверьте заголовки или "
                    "укажите --anon-column; НЕ отправляйте файл как есть (ревью wave5, C1).")
    if report.get("outside_rows"):
        warn.append(f"ВНИМАНИЕ: {report['outside_rows']} строк(и) ВНЕ таблицы (титул/шапка отчёта) "
                    "не обезличиваются — проверьте вручную (часто там ФИО).")
    if report.get("xlsx"):
        warn.append("ВНИМАНИЕ: имена листов и метаданные книги не обезличиваются (свойства книги очищены) — "
                    "просмотрите файл перед отправкой.")
    if report.get("empty_headers"):
        warn.append("ВНИМАНИЕ: есть колонки без заголовка — проверьте их вручную (могут содержать PII).")
    if report.get("merged_dropped"):
        warn.append(f"ВНИМАНИЕ: {report['merged_dropped']} объединённых диапазонов ниже усечённой части "
                    "удалены (openpyxl их не сдвигает) — проверьте файл.")
    print("\n".join(warn), file=sys.stderr, flush=True)


def _same_file(a: Path, b: Path) -> bool:
    """Совпадение файла-назначения с источником: тот же путь/симлинк ИЛИ hardlink (общий inode).

    resolve() НЕ ловит hardlink (две записи каталога на один inode): запись в такую «копию»
    обнулила бы исходную выписку. samefile сравнивает устройство+inode — ловит.
    """
    try:
        if a.resolve() == b.resolve():
            return True
    except OSError:
        pass
    try:
        return a.exists() and b.exists() and os.path.samefile(a, b)
    except OSError:
        return False


def run_anonymize(
    src: Path | str,
    dst: Path | str | None = None,
    *,
    max_rows: int = DEFAULT_ROWS,
    extra_columns: tuple[str, ...] | set[str] = (),
) -> int:
    """Обезличить файл (CSV/XLSX) и предупредить о неудаляемых датах/суммах. 0 — успех, 1 — ошибка."""
    utf8_stdout()
    if max_rows < 0:
        print("--rows не может быть отрицательным (0 — все строки)", file=sys.stderr, flush=True)
        return 1
    src = Path(src)
    raw = _read_source(src)
    if raw is None:
        return 1
    out = Path(dst) if dst else None
    if out is not None and _same_file(out, src):
        print("исходный и выходной файл совпадают — перезапись выписки запрещена; укажите другой путь",
              file=sys.stderr, flush=True)
        return 1
    try:
        data, report = _anonymize_payload(raw, src.suffix.lower(), set(extra_columns), max_rows)
    except (ValueError, OSError) as e:
        print(f"ошибка обезличивания: {e}", file=sys.stderr, flush=True)
        return 1
    except Exception as e:  # noqa: BLE001 — битый XLSX (не zip) — понятная ошибка, а не трейс
        print(f"не удалось разобрать файл: {e}", file=sys.stderr, flush=True)
        return 1
    if out is None:
        # авто-имя: расширение — по фактическому формату (XLSX под именем .csv → .anon.xlsx) — wave6 S5
        out = src.with_name(src.stem + ".anon" + (".xlsx" if report.get("xlsx") else src.suffix))
    if _same_file(out, src):
        # C3 (wave5, регрессия) + волна 5: guard действует и для авто-имени — симлинк/hardlink
        # v.anon.csv → v.csv, иначе затирал бы исходную выписку
        print("исходный и выходной файл совпадают — перезапись выписки запрещена; укажите другой путь",
              file=sys.stderr, flush=True)
        return 1
    if not _write_anonymized(data, out):
        return 1
    _print_report(report, out)
    return 0


def main(argv: list[str] | None = None) -> int:
    utf8_stdout()
    parser = argparse.ArgumentParser(description="Обезличить выписку CSV/XLSX для отправки образца")
    parser.add_argument("file", help="исходная выписка банка (CSV или XLSX)")
    parser.add_argument("output", nargs="?", default=None,
                        help="куда записать (по умолчанию <имя>.anon.csv рядом)")
    parser.add_argument("-o", "--out", dest="output_flag", default=None, metavar="ФАЙЛ",
                        help="то же, что позиционный аргумент (для скриптов)")
    parser.add_argument("--rows", type=int, default=DEFAULT_ROWS,
                        help=f"оставить первых N строк данных (0 — все; по умолчанию {DEFAULT_ROWS})")
    parser.add_argument("--anon-column", action="append", default=[],
                        help="доп. колонка для обезличивания (можно повторять)")
    args = parser.parse_args(argv)
    if args.rows < 0:
        parser.error("--rows не может быть отрицательным")
    if args.output and args.output_flag:
        parser.error("укажите выходной файл один раз: позиционным аргументом или -o/--out")
    return run_anonymize(args.file, args.output_flag or args.output,
                         max_rows=args.rows, extra_columns=args.anon_column)


if __name__ == "__main__":
    raise SystemExit(main())
